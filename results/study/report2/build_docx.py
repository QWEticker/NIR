"""Сборка REPORT2.docx из REPORT2.md с встраиванием рисунков и таблиц.

Используется python-docx (pandoc недоступен). Поддерживаются: заголовки #/##/###,
абзацы, маркированные и нумерованные списки, таблицы Markdown, блоки кода,
изображения ![alt](path), жирный **текст** и инлайн `код`.
"""
import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(HERE, "REPORT2.md")
OUT = os.path.join(HERE, "REPORT2.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Segoe UI"
style.font.size = Pt(11)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text):
    """Разбор инлайн-разметки: **жирный** и `код`."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            paragraph.add_run(part)


def resolve_img(path):
    for base in (HERE, os.path.join(HERE, "figures"), os.path.dirname(MD)):
        cand = os.path.normpath(os.path.join(base, path))
        if os.path.exists(cand):
            return cand
    cand = os.path.normpath(os.path.join(HERE, path))
    return cand if os.path.exists(cand) else None


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for i, c in enumerate(cells) if not re.match(r"^[-:\s|]+$", rows[i].strip().strip("|"))]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(cells):
        wr = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            wr[ci].text = ""
            p = wr[ci].paragraphs[0]
            add_runs(p, txt)
            if ri == 0:
                for run in p.runs:
                    run.bold = True


img_re = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
n = len(lines)
while i < n:
    line = lines[i]
    s = line.rstrip()

    if s.startswith("```"):
        i += 1
        buf = []
        while i < n and not lines[i].startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        p = doc.add_paragraph()
        r = p.add_run("\n".join(buf))
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        continue

    if s.strip() == "---":
        i += 1
        continue

    if s.startswith("# "):
        doc.add_heading(s[2:].strip(), level=0)
        i += 1
        continue
    if s.startswith("## "):
        doc.add_heading(s[3:].strip(), level=1)
        i += 1
        continue
    if s.startswith("### "):
        doc.add_heading(s[4:].strip(), level=2)
        i += 1
        continue

    m = img_re.match(s)
    if m:
        path = resolve_img(m.group(2))
        if path:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            try:
                run.add_picture(path, width=Inches(6.0))
            except Exception as e:
                p.add_run(f"[не удалось вставить {m.group(2)}: {e}]")
        else:
            doc.add_paragraph(f"[рисунок не найден: {m.group(2)}]")
        i += 1
        continue

    if s.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[-:\s|]+\|\s*$", lines[i + 1]):
        block = []
        while i < n and lines[i].lstrip().startswith("|"):
            block.append(lines[i])
            i += 1
        add_table(block)
        continue

    if re.match(r"^\s*[-*]\s+", s):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, re.sub(r"^\s*[-*]\s+", "", s))
        i += 1
        continue
    if re.match(r"^\s*\d+\.\s+", s):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\s*\d+\.\s+", "", s))
        i += 1
        continue

    if s.strip() == "":
        i += 1
        continue

    p = doc.add_paragraph()
    add_runs(p, s)
    i += 1

doc.save(OUT)
print(f"saved {OUT} ({os.path.getsize(OUT)} bytes)")
