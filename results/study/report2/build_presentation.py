"""Сборка PRESENTATION.docx из PRESENTATION.md (раскадровка + речь докладчика)."""
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(HERE, "PRESENTATION.md")
OUT = os.path.join(HERE, "PRESENTATION.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Segoe UI"
style.font.size = Pt(11)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            paragraph.add_run(part)


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for i, c in enumerate(cells)
             if not re.match(r"^[-:\s|]+$", rows[i].strip().strip("|"))]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(cells):
        wr = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            wr[ci].text = ""
            p = wr[ci].paragraphs[0]
            add_runs(p, txt)
            if ri == 0:
                for run in p.runs:
                    run.bold = True


lines = open(MD, encoding="utf-8").read().splitlines()
i, n = 0, len(lines)
while i < n:
    s = lines[i].rstrip()
    if s.strip() == "---":
        i += 1; continue
    if s.startswith("# "):
        doc.add_heading(s[2:].strip(), level=0); i += 1; continue
    if s.startswith("## "):
        doc.add_heading(s[3:].strip(), level=1); i += 1; continue
    if s.startswith("### "):
        doc.add_heading(s[4:].strip(), level=2); i += 1; continue
    if s.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[-:\s|]+\|\s*$", lines[i + 1]):
        block = []
        while i < n and lines[i].lstrip().startswith("|"):
            block.append(lines[i]); i += 1
        add_table(block); continue
    if re.match(r"^\s*[-*]\s+", s):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, re.sub(r"^\s*[-*]\s+", "", s)); i += 1; continue
    if re.match(r"^\s*\d+\.\s+", s):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\s*\d+\.\s+", "", s)); i += 1; continue
    if s.strip() == "":
        i += 1; continue
    p = doc.add_paragraph(); add_runs(p, s); i += 1

doc.save(OUT)
print(f"saved {OUT} ({os.path.getsize(OUT)} bytes)")
